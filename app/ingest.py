"""Passo 1 — Ingestão: ler ficheiros, partir em chunks, gerar embeddings, guardar na BD.

Este ficheiro orquestra os outros passos (chunking.py, embeddings.py, db.py).
"""

from pathlib import Path
from docx import Document as DocxDocument   
from app.chunking import chunk_text
from app.embeddings import embed_texts
from app.config import get_settings
from app.db import get_connection
from pgvector import Vector
from app.schemas import Document


def _read_docx(caminho: Path) -> str:
    d = DocxDocument(caminho)
    linhas = []
    for p in d.paragraphs:
        linhas.append(p.text)
    for tabela in d.tables:
        for linha in tabela.rows:
            linhas.append(" | ".join(c.text for c in linha.cells))
    return "\n".join(linhas)

def load_documents(docs_dir: Path) -> list[Document]:
    documentos = []
    for caminho in sorted(docs_dir.iterdir()):
        extensao = caminho.suffix.lower()
        if extensao in (".txt", ".md"):
            texto = caminho.read_text(encoding="utf-8")
        elif extensao == ".docx":
            texto = _read_docx(caminho)
        else:
            continue
        documentos.append(Document(source=caminho.name, title=caminho.stem, text=texto))
    return documentos


def ingest_directory(docs_dir: Path) -> int:
    """Pipeline completo de ingestão. Devolve o número de chunks guardados."""
    settings = get_settings()
    documentos = load_documents(docs_dir)
    total = 0

    with get_connection() as conn:
        for doc in documentos:
            conn.execute("DELETE FROM documents WHERE source = %s", (doc.source,))

            doc_id = conn.execute(
                "INSERT INTO documents (source, title) VALUES (%s, %s) RETURNING id",
                (doc.source, doc.title),
            ).fetchone()[0]
            chunks = chunk_text(doc.text, settings.chunk_size, settings.chunk_overlap)
            vetores = embed_texts(chunks)

            for i, (texto, vetor) in enumerate(zip(chunks, vetores)):
                conn.execute(
                    "INSERT INTO chunks (document_id, chunk_index, content, embedding) "
                    "VALUES (%s, %s, %s, %s)",
                    (doc_id, i, texto, Vector(vetor)),)                
                total += 1

    return total


if __name__ == "__main__":
    for doc in load_documents(Path("data/docs")):
        print(doc.source, len(doc.text))